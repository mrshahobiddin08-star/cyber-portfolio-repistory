import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

def create_cyber_portfolio():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts

    COLOR_DARK_BLACK = RGBColor(11, 11, 15)
    COLOR_NEON_CRIMSON = RGBColor(220, 20, 60)
    COLOR_WHITE = RGBColor(245, 245, 247)

    # SLAYD 1
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_DARK_BLACK
    bg1.line.fill.background()

    tx_box1 = slide1.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(6.0), Inches(4.5))
    tf1 = tx_box1.text_frame
    tf1.word_wrap = True
    
    p1 = tf1.paragraphs
    p1.text = "ARCHITECTING THE FUTURE"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_NEON_CRIMSON
    p1.font.name = "Impact"

    p2 = tf1.add_paragraph()
    p2.text = "\nFusing advanced system core programming with cutting-edge Artificial Intelligence networks to shift the limits of global technology."
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_WHITE
    p2.font.name = "Consolas"

    image_path1 = "assets/hacker_me.jpg"
    if os.path.exists(image_path1):
        slide1.shapes.add_picture(image_path1, Inches(6.8), Inches(0.8), width=Inches(5.8), height=Inches(5.9))

    # SLAYD 2
    slide2 = prs.slides.add_slide(blank_layout)
    bg2 = slide2.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = COLOR_DARK_BLACK
    bg2.line.fill.background()

    image_path2 = "assets/future_newspaper.jpg"
    if os.path.exists(image_path2):
        slide2.shapes.add_picture(image_path2, Inches(0.6), Inches(0.8), width=Inches(5.8), height=Inches(5.9))

    tx_box2 = slide2.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(6.0), Inches(4.5))
    tf2 = tx_box2.text_frame
    tf2.word_wrap = True

    p3 = tf2.paragraphs
    p3.text = "GLOBAL RECOGNITION"
    p3.font.size = Pt(36)
    p3.font.bold = True
    p3.font.color.rgb = COLOR_NEON_CRIMSON
    p3.font.name = "Impact"

    p4 = tf2.add_paragraph()
    p4.text = "\nBreaking headlines as a world-renowned Economist and AI Engineer. Successfully merging algorithmic market economic strategies with deep neural programming to lead the next multi-trillion dollar digital era."
    p4.font.size = Pt(16)
    p4.font.color.rgb = COLOR_WHITE
    p4.font.name = "Arial"

    prs.save("Cyber_AI_Visionary_Portfolio.pptx")
    print("Prezentatsiya muvaffaqiyatli yaratildi!")

if __name__ == "__main__":
    create_cyber_portfolio()
